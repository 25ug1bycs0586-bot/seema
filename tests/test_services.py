from io import BytesIO

from docx import Document

from services.learning_recommender import build_learning_plan
from services.resume_parser import ResumeParseError, extract_resume_text
from services.skill_analyzer import analyze_resume, detect_skills


def test_detect_skills_distinguishes_mentions_from_nearby_evidence():
    skills = detect_skills("SKILLS\nPython\nPROJECTS\nBuilt a small Python API.")
    by_name = {item["name"]: item for item in skills}

    assert by_name["python"]["evidence_level"] == "supported"

    isolated = {item["name"]: item for item in detect_skills("Skills: SQL, Git.")}
    assert isolated["sql"]["evidence_level"] == "mentioned"
    assert isolated["git"]["evidence_level"] == "mentioned"


def test_analysis_scores_are_transparent_and_missing_is_not_absence():
    report = analyze_resume(
        "Skills: Python, SQL. Projects: Built a small Python tool.",
        target_role="Python Developer",
    )

    assert report["job_match"]["required_count"] == 5
    assert report["job_match"]["score"] == 40
    assert {item["skill"] for item in report["missing_skills"]} == {"flask", "django", "git"}
    assert all("does not establish" in item["note"] for item in report["missing_skills"])
    assert report["quality"]["maximum"] == 100
    assert "not a hiring" in report["job_match"]["method"]


def test_learning_plan_has_unique_trackable_tasks_and_real_resources():
    plan = build_learning_plan(["python"])
    tasks = [task for phase in plan for task in phase["tasks"]]

    assert len(plan) == 4
    assert len({task["title"] for task in tasks}) == 4
    assert all(task["exercise"] and task["goal"] for task in tasks)
    assert all(task["resource_url"].startswith("https://") for task in tasks)


def test_docx_text_is_extracted_and_invalid_pdf_is_rejected():
    document = Document()
    document.add_paragraph("Education: Computer Science")
    document.add_paragraph("Projects: Python application")
    output = BytesIO()
    document.save(output)

    text = extract_resume_text("resume.docx", output.getvalue())
    assert "Computer Science" in text
    assert "Python application" in text

    try:
        extract_resume_text("resume.pdf", b"not a pdf")
    except ResumeParseError as error:
        assert "valid PDF" in str(error)
    else:
        raise AssertionError("Invalid PDF content should not be accepted.")
