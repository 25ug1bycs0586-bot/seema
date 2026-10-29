from io import BytesIO
from zipfile import BadZipFile, ZipFile

import pymupdf
from docx import Document


class ResumeParseError(ValueError):
    """Raised when an uploaded document cannot be safely parsed."""


def extract_resume_text(filename, content):
    if not filename or not content:
        raise ResumeParseError("Choose a non-empty PDF or DOCX resume.")
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    try:
        if extension == "pdf":
            if not content.startswith(b"%PDF-"):
                raise ResumeParseError("The selected file is not a valid PDF document.")
            with pymupdf.open(stream=content, filetype="pdf") as document:
                text = "\n".join(page.get_text("text") for page in document)
            if len(text.strip()) < 30:
                raise ResumeParseError(
                    "This PDF contains little extractable text. It may be a scanned image; OCR is not built in."
                )
            return text
        if extension == "docx":
            with ZipFile(BytesIO(content)) as archive:
                if "word/document.xml" not in archive.namelist():
                    raise ResumeParseError("The DOCX file is missing its document content.")
            document = Document(BytesIO(content))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            text += "\n" + "\n".join(
                " | ".join(cell.text for cell in row.cells)
                for table in document.tables
                for row in table.rows
            )
            if not text.strip():
                raise ResumeParseError("No readable text was found in the DOCX file.")
            return text
    except (pymupdf.FileDataError, pymupdf.EmptyFileError, BadZipFile, OSError) as error:
        raise ResumeParseError("The selected document is unreadable or damaged.") from error
    except Exception as error:
        if isinstance(error, ResumeParseError):
            raise
        raise ResumeParseError("The selected document could not be parsed.") from error
    raise ResumeParseError("Only PDF and DOCX resumes are supported.")
